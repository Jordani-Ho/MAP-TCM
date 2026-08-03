from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "中医智能体产品需求说明书_v0.1.docx"

FONT = "SimSun"
ACCENT = "2E74B5"
DARK = "1F4D78"
MUTED = "666666"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"


def set_run_font(run, size=11, bold=None, color=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    """Use fixed DXA geometry so Word and LibreOffice agree on table width."""
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=MUTED)


def add_para(doc, text="", size=11, bold=False, color=None, italic=False, after=6, before=0, align=None, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_rich_para(doc, chunks, after=6, before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    for text, options in chunks:
        run = p.add_run(text)
        set_run_font(run, **options)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_list(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        run = p.add_run(item)
        set_run_font(run, size=11)


def add_numbered_list(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        run = p.add_run(item)
        set_run_font(run, size=11)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, text in zip(header.cells, headers):
        shade(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        set_run_font(r, size=10, bold=True, color=DARK)
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(value)
            set_run_font(r, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, CALLOUT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    r1 = p.add_run(label + "  ")
    set_run_font(r1, size=10.5, bold=True, color=DARK)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for level, size, color, before, after in (
        (1, 16, ACCENT, 16, 8),
        (2, 13, ACCENT, 12, 6),
        (3, 12, DARK, 8, 4),
    ):
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(11)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    r = header.add_run("中医智能体产品需求说明书 | 内部讨论稿")
    set_run_font(r, size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    return doc


def add_title_block(doc):
    add_para(doc, "产品需求说明书", size=11, bold=True, color=ACCENT, after=6)
    add_para(doc, "中医智能体：医生的专业助理与徒弟", size=24, bold=True, color=DARK, after=8)
    add_para(doc, "将跨境远程问诊中的多媒体材料、辨证结论与治疗方案，沉淀为医生可复核的长期健康病历", size=13, color=MUTED, after=20)
    rows = [
        ("版本", "V0.1 - 需求澄清稿"),
        ("日期", "2026年7月"),
        ("状态", "待使用者与中医医生共同确认"),
        ("首要目标", "先对齐真实服务流程，再进入原型和动画短视频制作"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1800, 7560])
    for row, (label, value) in zip(table.rows, rows):
        shade(row.cells[0], LIGHT_BLUE)
        for cell in row.cells:
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
        r1 = row.cells[0].paragraphs[0].add_run(label)
        set_run_font(r1, size=10, bold=True, color=DARK)
        r2 = row.cells[1].paragraphs[0].add_run(value)
        set_run_font(r2, size=10)
    add_para(doc, "", after=8)
    add_callout(doc, "本文件的产品边界", "AI负责采集、整理、提示与协作；医生负责辨证、治疗决策和对外沟通。产品不会自主诊断、开方或替代紧急医疗服务。")
    add_heading(doc, "一页摘要", 1)
    add_para(doc, "这是一个面向中医医生、居家护理人员和长期患者的“病历与协作助理”。它把散落在微信中的文字、语音、舌象、患处照片、视频和医生回复，整理成按时间延续的结构化中医病历草稿；由医生确认辨证与治疗方案后，系统再承担预约、用药/治疗提醒、随访和反馈归档。")
    add_para(doc, "首个可验证版本不追求像医生一样回答问题，而是先做到：一次问诊结束后，医生用更少时间得到一份完整、可编辑、可追溯的病历；护理人员和患者知道下一步该做什么；下一次问诊不必从零回忆。")


def build(doc):
    add_title_block(doc)

    add_heading(doc, "1. 项目背景与问题", 1)
    add_para(doc, "目标使用者是一位居住在加州湾区、长期照顾老人的居家护理工作人员。多年中，用户通过微信向中国的中医朋友描述便秘、失眠、夜尿多等长期不适，并发送舌象、患处照片、语音与文字；医生依据这些材料给出辨证和治疗方案，包括中药与针灸建议。有效信息与闲聊、重复信息混在一起，长期积累后已难以回看、比较和复用。")
    add_heading(doc, "核心痛点", 2)
    add_list(doc, [
        "问诊材料分散：同一次问诊的语音、照片、视频、文字、医生回复散落在聊天窗口，无法形成完整病历。",
        "信息不连续：症状变化、既往方药、舌象变化和疗效反馈缺少时间线，医生每次都要重新问。",
        "医生整理负担高：医生的辨证和治疗方案常以语音或简短聊天回复出现，难以沉淀为可检索、可复核的临床记录。",
        "照护协作容易断：预约、复诊、服药/治疗、观察反馈常由护理人员靠记忆执行，老人和家属之间的信息不对称。",
        "历史资料价值未被释放：多年微信记录既有临床线索，也有噪声；现阶段缺少安全、可控的整理机制。",
    ])

    add_heading(doc, "2. 产品定位与目标", 1)
    add_callout(doc, "一句话定位", "一名住在手机里的中医医生助理：听得懂患者和护理人员的描述，看得见问诊材料，整理得出病历草稿，提醒得住预约和随访，但所有辨证与治疗结论都由医生确认。")
    add_heading(doc, "2.1 产品角色", 2)
    add_table(doc, ["角色", "主要任务", "产品获得的价值"], [
        ("中医医生", "查看病历草稿，补充四诊信息，确认辨证与治疗方案，安排复诊。", "减少重复整理，把时间留给判断、沟通与治疗。"),
        ("居家护理人员", "代为记录症状、上传材料、执行提醒、提交疗效与异常反馈。", "形成清楚的照护清单，减少漏记、漏问、漏随访。"),
        ("患者/老人", "表达感受，查看经医生确认的计划，接收合适的提醒。", "不用反复回忆既往信息，获得连续的照护体验。"),
        ("家属/协作者", "在授权范围内了解计划与反馈。", "知道照护进展，必要时共同配合。"),
    ], [1700, 4000, 3660])
    add_heading(doc, "2.2 首期目标", 2)
    add_numbered_list(doc, [
        "把一次新问诊中的多媒体材料汇集到一个病例，并按固定结构生成“待医生确认”的病历草稿。",
        "让医生能快速补充、修改和确认辨证与治疗方案，且保留每次修改的责任归属和时间记录。",
        "让护理人员/患者获得明确的预约、随访和反馈任务，并能把结果回流到下一次问诊。",
        "建立按患者维度的长期健康时间线，为后续个体健康模型和医生个人知识整理打下数据基础。",
    ])
    add_heading(doc, "2.3 非目标", 2)
    add_list(doc, [
        "不由AI自主给出诊断、辨证结论、处方、针灸方案或剂量。",
        "不在第一阶段直接读取、抓取或批量同步微信私聊记录；历史材料以用户主动粘贴、上传、转发或人工归档为主。",
        "不把平台做成公开问诊市场、药品交易平台或社交内容平台。",
        "不处理紧急医疗服务；产品只可展示经专业人士配置的提示与升级处理流程。",
    ])

    add_heading(doc, "3. 明确假设与待确认事项", 1)
    add_heading(doc, "3.1 当前假设", 2)
    add_table(doc, ["编号", "假设", "原因与影响"], [
        ("A1", "首批用户是你、你照护的老人，以及少量与你已有信任关系的中医医生。", "先做小范围闭环，避免在未验证前处理公开获客和大规模运营。"),
        ("A2", "首期使用中文为主，并允许夹带英文药名、检查结果或本地就医信息。", "贴合中国中医医生与湾区护理场景；双语界面可作为后续需求。"),
        ("A3", "首期历史资料导入采用手动选择与人工确认，不依赖微信的直接数据接口。", "避免一开始被平台授权、隐私和数据完整性问题卡住。"),
        ("A4", "AI生成内容全部标记为“草稿/待确认”，未经医生确认不得向患者呈现为诊疗结论。", "保证临床责任清晰，也让产品定位稳定。"),
        ("A5", "产品形态先不锁死为小程序、聊天机器人或本地智能体；先以同一业务内核验证，再选最适合的入口。", "视频和原型可以先验证价值，不必过早押注技术载体。"),
    ], [700, 3860, 4800])
    add_heading(doc, "3.2 需要在需求评审中确认的问题", 2)
    add_numbered_list(doc, [
        "首个可用版本的主要入口是什么：微信小程序、独立手机应用、聊天机器人，还是医生端网页加患者端轻入口？",
        "参与的中医医生是否愿意每天在系统内复核病历草稿？他们希望用手机、电脑还是微信消息接收待办？",
        "历史微信记录优先整理哪些患者、哪些时间段、哪些病种？是否需要把原始聊天截图作为病历附件长期保留？",
        "医生确认后的治疗方案需要呈现到什么颗粒度：仅展示医嘱摘要，还是包括方名、药味、剂量、针灸穴位、疗程及注意事项？",
        "老人本人、家属、护理人员各自能看到哪些内容？谁有权上传、编辑、撤回或删除资料？",
        "产品最初是否只做记录与协作，还是同时接入预约收费、药房/针灸师协作等能力？",
    ])

    add_heading(doc, "4. 用户与典型场景", 1)
    add_heading(doc, "4.1 关键用户画像", 2)
    add_table(doc, ["用户", "使用环境", "关键诉求", "主要顾虑"], [
        ("护理人员", "在家中照护，常用手机，时间碎片化。", "快速说清情况、留住照片和语音、收到下一步提醒。", "不想填写复杂表格；担心遗漏和隐私泄露。"),
        ("中医医生", "通过微信/手机处理跨时区咨询，回复多为语音或简讯。", "快速掌握病程，保有临床判断权，少做文书。", "不愿被AI替代或被不完整资料误导。"),
        ("老人/患者", "可能不擅长输入，症状表达不稳定。", "简单表达感受，得到清楚、可执行的医生确认计划。", "怕操作复杂，也需要家属协助。"),
    ], [1500, 2300, 3000, 2560])
    add_heading(doc, "4.2 主流程：从一次问诊到下一次随访", 2)
    add_numbered_list(doc, [
        "护理人员创建一次问诊，选择患者，并上传/录入文字、语音、舌象、患处照片、短视频和已知聊天回复。",
        "助理转写语音、为材料标注时间和来源，提示缺失信息，例如饮食、饮水、二便、睡眠、情绪、体温、既往用药等。",
        "助理从原始材料中提取事实并生成结构化病历草稿，明确区分“患者/护理人员陈述”“附件材料”“AI整理”“医生确认”。",
        "中医医生查看草稿，补充望闻问切信息，填写或语音输入辨证与治疗方案，并确认可见范围。",
        "系统把经确认的计划转化为预约、复诊、观察、反馈等任务；护理人员按期回填结果。",
        "下次问诊时，助理自动呈现上次计划、已完成情况、疗效反馈与待追问事项，形成连续闭环。",
    ])
    add_callout(doc, "关键设计原则", "保留原始证据，但不把原始聊天当作病历。每一条结构化记录应可追溯回原文、语音转写、照片或医生确认动作。")

    add_heading(doc, "5. MVP 功能需求", 1)
    add_heading(doc, "5.1 模块总览", 2)
    add_table(doc, ["模块", "MVP能力", "医生确认要求"], [
        ("患者与授权", "建立患者档案、照护关系、授权范围与联系人。", "患者/授权照护者完成授权；医生只访问被授权患者。"),
        ("多媒体问诊盒", "文字、语音、图片、视频、文件上传；语音转写；材料时间线。", "无需医生逐条确认，但可标记关键证据。"),
        ("问诊引导", "按固定问题补齐饮食、饮水、二便、睡眠等信息；支持跳过与稍后填写。", "问诊问题模板可由医生配置。"),
        ("AI病历草稿", "提取主诉、现病史、四诊材料、既往史、反馈、待追问项。", "辨证、诊断性措辞和治疗建议必须待医生确认。"),
        ("医生工作台", "查看原始材料与草稿、编辑、语音录入、确认/退回、生成患者可见摘要。", "医生是最终提交者。"),
        ("预约与随访", "预约请求、复诊日期、提醒、反馈表、未完成任务列表。", "提醒内容和频率可由医生/机构配置。"),
        ("健康时间线", "按日期查看症状、附件、方案、反馈、预约。", "医生确认的记录与AI草稿视觉上必须明显区分。"),
    ], [1700, 4500, 3160])
    add_heading(doc, "5.2 多媒体问诊盒", 2)
    add_list(doc, [
        "支持一次问诊下的文字、语音、图片、短视频、文档和聊天截图，并保存材料来源、上传者、时间、患者归属。",
        "语音转写后保留原音频，可人工修正转写内容；转写不能覆盖原始材料。",
        "图片/视频需要用于提示拍摄时间、部位、光线和角度，但不把图像自动判定为诊断结论。",
        "对历史微信材料提供“选择后归档”流程：先选中一段材料，再指定患者、问诊日期与来源；不做自动全量导入。",
        "上传前显示授权提示，允许为包含他人信息的材料设置更严格可见范围。",
    ])
    add_heading(doc, "5.3 结构化中医病历", 2)
    add_para(doc, "病历结构以长期追踪和医生复核为核心。每个字段均记录来源类型、采集时间、录入人和确认状态。")
    add_table(doc, ["病历区块", "核心字段", "来源/状态"], [
        ("基础信息", "患者身份、性别、年龄、既往病史、过敏/禁忌、就诊时间、地点、照护关系。", "患者/护理人员填写；可由医生修正。"),
        ("患者与护理人员陈述", "主诉、症状起止、程度、诱因、饮食饮水、二便、睡眠、情绪、既往处理。", "原始输入 + AI整理；保留可追溯链接。"),
        ("望闻问切材料", "舌象照片、患处照片、视频、气味/声音描述、医生问诊记录、脉诊等。", "附件或医生录入；医生可标记为关键。"),
        ("医生辨证", "辨证思路、证候、治则治法、需排除或需进一步询问事项。", "仅医生确认后生效。"),
        ("治疗方案", "中药/方剂、针灸、饮食调护、外治、生活建议、频次、复诊计划、注意事项。", "仅医生确认后向患者端发布。"),
        ("疗效与随访", "执行情况、症状变化、不良/异常反馈、患者评价、下次调整依据。", "患者/护理人员提交；医生可评阅。"),
    ], [1700, 4820, 2840])
    add_heading(doc, "5.4 AI整理规则", 2)
    add_list(doc, [
        "只提取材料中已出现的事实；不确定内容要标明“待确认”，不能把猜测写成患者陈述。",
        "将相同或相近症状按时间归并，但原始表达与附件必须可回看。",
        "将医生语音/文字中的明确指令提取为草稿字段，并要求医生确认后才生成面向患者的任务。",
        "在病历中显著标注AI生成段落、人工修改段落和医生确认时间。",
        "在医生未确认时，不生成“处方已开立”“治疗已建议”等对外表述。",
    ])

    add_heading(doc, "6. 预约、随访与照护协作", 1)
    add_heading(doc, "6.1 预约", 2)
    add_list(doc, [
        "患者或护理人员提交预约意向：可选时段、问题摘要、是否需要视频问诊、是否有新增照片/视频。",
        "医生端接受、改期或拒绝，并输出明确的时间、时区、沟通方式和会前准备事项。",
        "系统用患者本地时区展示提醒，同时保留医生端时区，避免中美沟通中的时间误差。",
    ])
    add_heading(doc, "6.2 随访", 2)
    add_table(doc, ["触发点", "系统动作", "预期输出"], [
        ("医生确认方案", "生成医生配置的随访任务与到期时间。", "护理人员待办与患者端提醒。"),
        ("护理人员提交反馈", "归档反馈，按模板汇总变化与未回答项。", "医生可快速评阅的随访摘要。"),
        ("到期未反馈", "在授权与配置范围内提醒护理人员，不作医疗结论。", "补充信息或重新预约。"),
        ("下一次问诊创建", "带入上次方案、完成状态、关键变化和待追问事项。", "连续病程视图。"),
    ], [1800, 4500, 3060])
    add_heading(doc, "6.3 家庭照护模式", 2)
    add_para(doc, "一个老人可能由护理人员、子女和医生共同参与。系统需支持“按角色、按病例、按内容类型”授权：例如家属可看到预约和医生确认摘要，但不能查看全部照片；护理人员可以上传日常反馈，但不能修改医生方案。授权变更、导出、撤回和删除均要留痕。")

    add_heading(doc, "7. 产品形态建议", 1)
    add_para(doc, "本产品的本质是一个“记录与协作内核”，入口可以变化。首期应依据医生实际工作习惯选择，而不是由技术偏好决定。")
    add_table(doc, ["形态", "适合之处", "主要限制", "建议"], [
        ("微信小程序", "护理人员和患者进入门槛低，适合上传、提醒和查看时间线。", "历史聊天材料与平台能力受限，医生复杂编辑体验一般。", "适合作为患者/护理端轻入口。"),
        ("聊天机器人", "符合医生和护理人员已有的对话式工作习惯，采集体验自然。", "复杂病历编辑、权限和附件管理需要额外界面。", "适合作为问诊引导和待办提醒入口。"),
        ("本地手机智能体", "可更深地辅助整理本机材料、语音输入和提醒。", "隐私、系统权限、跨设备同步和维护成本更高。", "适合在验证后作为高隐私/高频使用版本。"),
        ("医生网页工作台", "适合长病历审核、批量随访和检索。", "患者端不宜作为主要入口。", "建议与轻入口配合，作为医生复核主界面。"),
    ], [1500, 2900, 2700, 2260])
    add_callout(doc, "阶段性建议", "先用“患者/护理端轻入口 + 医生端复核工作台”的双端思路做闭环原型；具体是小程序还是聊天入口，可由参与医生的日常习惯决定。动画短视频则用于验证大家是否一眼理解这名“徒弟”的价值。")

    add_heading(doc, "8. 数据、隐私与安全要求", 1)
    add_para(doc, "本产品将处理高度敏感的健康信息、照片、视频和语音。以下为产品要求，不构成法律意见；在面向加州患者和跨境医生正式提供服务前，需由具备相应资质的法律与合规专业人士完成具体评估。")
    add_heading(doc, "8.1 必备控制", 2)
    add_list(doc, [
        "最小权限：按患者、角色和资料类型授权，默认不共享；医生只能访问被授权病例。",
        "明确同意：在首次建档、上传敏感照片/视频、邀请家属、跨境传输或AI处理前，提供可理解的授权说明与撤回入口。",
        "安全留存：传输与存储加密、强身份验证、设备锁定、会话超时、访问与导出审计日志。",
        "数据控制：支持查看、导出、更正申请、删除/留存策略；原始材料与AI整理结果分别管理。",
        "模型使用边界：默认不得把可识别健康信息用于训练通用模型；若使用第三方AI服务，需明确数据处理协议和地域策略。",
        "事件处理：建立异常访问、泄露、误发资料的发现、处置、通知与复盘流程。",
    ])
    add_heading(doc, "8.2 远程服务与跨境边界", 2)
    add_para(doc, "产品在设计上应区分“健康资料整理工具”和“向位于加州的患者提供远程专业服务”两种模式。后者可能涉及执照、职业监管、隐私、病历留存、药品与支付等要求，不能因医生位于中国或沟通使用微信而自动豁免。加州医疗委员会指出，向加州患者使用技术开展医疗执业需由持有有效加州执照的医师进行，并遵守州和联邦隐私法律；加州针灸委员会亦对向加州患者提供远程针灸服务发布了专门说明。正式上线前应逐项核验服务主体与服务范围。")
    add_heading(doc, "8.3 参考资料（供合规评审起点使用）", 2)
    add_list(doc, [
        "美国卫生与公众服务部（HHS）：HIPAA and Telehealth：https://www.hhs.gov/hipaa/for-professionals/special-topics/telehealth/index.html",
        "HHS Telehealth：Privacy laws and policy guidance：https://telehealth.hhs.gov/providers/best-practice-guides/privacy-and-security-telehealth/privacy-laws-and-policy-guidance",
        "加州医疗委员会：Telehealth：https://mbc.ca.gov/Resources/Medical-Resources/telehealth.aspx/",
        "加州针灸委员会：Telehealth and Acupuncture：https://www.acupuncture.ca.gov/pubs_forms/telehealth_statement.pdf",
    ])

    add_heading(doc, "9. MVP 验收标准", 1)
    add_table(doc, ["编号", "验收结果", "验证方式"], [
        ("AC-01", "护理人员可以为指定患者创建一次问诊，并上传至少文字、语音、图片三类材料。", "完成一例模拟问诊并检查材料归属、时间与来源。"),
        ("AC-02", "系统能将问诊材料整理为六个病历区块，并对缺失内容列出待追问项。", "由医生比较原始材料和草稿，确认没有把推测写成事实。"),
        ("AC-03", "医生可以编辑并确认辨证和治疗方案；确认前，患者端不可把草稿当作医嘱查看。", "用确认前/后两种状态进行权限与界面测试。"),
        ("AC-04", "医生确认后可生成预约/随访任务，护理人员可提交完成情况和疗效反馈。", "运行一次“方案-提醒-反馈-下次问诊”的闭环。"),
        ("AC-05", "同一患者的两次以上问诊可在时间线上查看，并能回到原始附件或医生确认记录。", "选取历史案例完成检索与回溯测试。"),
        ("AC-06", "用户能查看并管理谁可访问患者资料，系统记录关键访问和导出操作。", "进行授权、撤销与审计演示。"),
    ], [850, 5400, 3110])

    add_heading(doc, "10. 分阶段实施建议", 1)
    add_table(doc, ["阶段", "目标", "交付物", "是否进入下一阶段的判断"], [
        ("0. 需求确认", "与护理人员和2-3位医生校准真实工作流、字段和责任边界。", "本说明书定稿、字段清单、样例病例。", "每位参与者都能确认流程不改变其工作习惯的核心部分。"),
        ("1. 概念表达", "把产品价值说清楚、看得见。", "60-90秒动画短视频脚本、分镜和视觉稿。", "观众能复述“它整理、提醒，但不替医生做决定”。"),
        ("2. 可点击原型", "验证采集、草稿、医生确认、随访四步是否顺手。", "手机端和医生端交互原型。", "完成5-10个真实或脱敏案例走查。"),
        ("3. 封闭MVP", "在少量被邀请用户中跑通真实闭环。", "可用产品、授权与审计基础能力。", "医生复核耗时下降，护理人员愿意持续填写反馈。"),
        ("4. 合规与扩展", "根据服务地区、主体和业务范围完善合规与规模化能力。", "合规评估、数据策略、运营方案。", "在明确责任与风险控制后再扩大用户。"),
    ], [1450, 2500, 3000, 2410])

    add_heading(doc, "11. 动画短视频制作简报（需求确认后执行）", 1)
    add_para(doc, "短视频不是宣传“AI能看病”，而是让医生、护理人员和家属在90秒内看懂它如何把混乱问诊变成连续照护。建议时长60-90秒，风格温暖、真实、克制，主角是护理人员、老人和医生，而非炫技的机器人。")
    add_table(doc, ["镜头", "画面", "信息点"], [
        ("1", "湾区家中，护理人员照顾老人，手机里有微信语音、舌象照片和长聊天。", "信息很多，但没有病历。"),
        ("2", "护理人员把一次问诊的文字、语音、照片和视频放入“徒弟”。", "多媒体采集，不增加复杂填写。"),
        ("3", "助手整理出“待确认病历”，并温和提示还缺少的观察信息。", "AI整理事实，不下结论。"),
        ("4", "中国的中医医生在工作台上查看原始材料，补充辨证和方案并确认。", "医生保有判断与责任。"),
        ("5", "护理人员收到预约与随访提醒，回填执行和症状变化。", "照护变得连续。"),
        ("6", "时间线串起多次问诊，医生看到变化，老人和家属安心。", "这名徒弟让长期经验沉淀下来。"),
    ], [700, 5300, 3360])
    add_callout(doc, "视频的最后一句", "“我不替医生看病。我把每一次认真看病，留下来，接得上。”")

    add_heading(doc, "12. 下一次需求评审建议", 1)
    add_numbered_list(doc, [
        "从现有微信记录中选择1-2个已脱敏的典型案例，按本病历结构做一次人工整理，检验字段是否够用。",
        "邀请至少一位中医医生逐段审阅“医生工作台”和“病历字段”，确认其愿意确认哪些内容、拒绝哪些自动化。",
        "确定首期产品入口和数据范围：只做新问诊，还是包含一小批历史记录归档。",
        "在开始技术开发和动画制作前，确认产品名称、服务主体、目标地区、语言、隐私同意方式与医生端使用方式。",
    ])


def main():
    doc = setup_document()
    build(doc)
    doc.core_properties.title = "中医智能体产品需求说明书 V0.1"
    doc.core_properties.subject = "中医医生专业助理与徒弟的产品需求说明书"
    doc.core_properties.author = "项目团队"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
